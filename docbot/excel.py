"""Распознанный документ → Excel-файл в памяти (бот присылает его в чат)."""

from __future__ import annotations

import io

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill

from docbot.schemas import Document

HEADER_FILL = PatternFill("solid", fgColor="1F4E79")
WARN_FILL = PatternFill("solid", fgColor="FFF2CC")


def make_excel(doc: Document, issues: list[str] | None = None) -> io.BytesIO:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Документ"

    other = doc.customer if doc.is_act else doc.buyer
    header = [
        ("Тип документа", "Акт выполненных работ" if doc.is_act else "Накладная"),
        ("Номер", doc.number),
        ("Дата", doc.date),
        ("Исполнитель" if doc.is_act else "Поставщик", doc.counterparty.name),
        ("ИИН/БИН", doc.counterparty.bin),
        ("Заказчик" if doc.is_act else "Получатель", other.name if other else ""),
    ]
    row = 1
    for key, value in header:
        ws.cell(row=row, column=1, value=key).font = Font(bold=True)
        ws.cell(row=row, column=2, value=value)
        row += 1

    row += 1
    for col, name in enumerate(["Описание" if doc.is_act else "Товар", "Кол-во", "Ед.", "Цена", "Сумма"], 1):
        cell = ws.cell(row=row, column=col, value=name)
        cell.fill = HEADER_FILL
        cell.font = Font(bold=True, color="FFFFFF")
        cell.alignment = Alignment(horizontal="center")
    row += 1

    for line in doc.lines:
        for col, value in enumerate([line.title, line.quantity, line.unit, line.price, line.amount], 1):
            ws.cell(row=row, column=col, value=value)
        row += 1

    row += 1
    for label, value in [("Итого", doc.total), ("НДС", doc.vat), ("Итого с НДС", doc.total_with_vat)]:
        ws.cell(row=row, column=1, value=label).font = Font(bold=True)
        ws.cell(row=row, column=5, value=value)
        row += 1

    if issues:
        row += 1
        ws.cell(row=row, column=1, value="Требует проверки").font = Font(bold=True)
        for issue in issues:
            row += 1
            ws.cell(row=row, column=1, value=issue).fill = WARN_FILL

    ws.column_dimensions["A"].width = 45
    for col in "BCDE":
        ws.column_dimensions[col].width = 15

    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return buffer
